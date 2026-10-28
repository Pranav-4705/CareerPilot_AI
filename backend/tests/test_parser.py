import io
import os
import fitz
import docx
from fastapi import status
from app.models.resume import Resume
from app.models.user import User


def get_auth_token(client, email="parser.candidate@careerpilot.ai", name="Parser User"):
    """Helper to register and login a test user to get a Bearer token."""
    reg_payload = {
        "name": name,
        "email": email,
        "password": "Password123!",
    }
    client.post("/api/v1/auth/register", json=reg_payload)

    login_res = client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": "Password123!"},
    )
    return login_res.json()["access_token"]


def create_mock_pdf_bytes(text: str = "John Doe\nSenior Software Engineer\nSkills: Python, FastAPI, PostgreSQL") -> bytes:
    """Generates a valid single-page PDF with the specified text."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_multipage_pdf_bytes() -> bytes:
    """Generates a valid 3-page PDF with distinct text on each page."""
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 72), "Page 1: Senior Software Engineer with 6 years experience.")
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Page 2: Projects include CareerPilot AI and Scalable Cloud Platforms.")
    page3 = doc.new_page()
    page3.insert_text((50, 72), "Page 3: B.S. in Computer Science and AWS Certification.")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_mock_docx_bytes() -> bytes:
    """Generates a valid DOCX document with paragraphs and tables."""
    doc = docx.Document()
    doc.add_heading("Jane Candidate - Full Stack Engineer", 0)
    doc.add_paragraph("Summary: Experienced developer proficient in TypeScript, Python, and PostgreSQL.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Proficiency"
    table.cell(1, 0).text = "FastAPI"
    table.cell(1, 1).text = "Expert"

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def test_parse_single_page_pdf(client, db_session):
    """Test extracting text from a single-page PDF resume."""
    token = get_auth_token(client, email="pdf.parse@example.com")
    pdf_bytes = create_mock_pdf_bytes("Alice Developer\nBackend Engineer\nProficient in Python and FastAPI")

    # 1. Upload
    upload_res = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("alice_resume.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert upload_res.status_code == status.HTTP_201_CREATED
    resume_id = upload_res.json()["id"]

    # 2. Parse
    parse_res = client.post(
        "/api/v1/resumes/{}/parse".format(resume_id),
        headers={"Authorization": f"Bearer {token}"},
    )
    assert parse_res.status_code == status.HTTP_200_OK
    parse_data = parse_res.json()
    assert parse_data["id"] == resume_id
    assert parse_data["status"] == "parsed"
    assert parse_data["extracted_text_length"] > 20

    # 3. Verify Get details & DB persistence
    get_res = client.get(
        "/api/v1/resumes/{}".format(resume_id),
        headers={"Authorization": f"Bearer {token}"},
    )
    assert get_res.status_code == status.HTTP_200_OK
    resume_data = get_res.json()
    assert "Alice Developer" in resume_data["extracted_text"]
    assert "Backend Engineer" in resume_data["extracted_text"]

    db_record = db_session.query(Resume).filter(Resume.id == resume_id).first()
    assert db_record.extracted_text is not None
    assert "Alice Developer" in db_record.extracted_text


def test_parse_multipage_pdf(client):
    """Test extracting text from a multi-page PDF resume."""
    token = get_auth_token(client, email="multipage@example.com")
    pdf_bytes = create_multipage_pdf_bytes()

    upload_res = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("multipage_cv.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
        headers={"Authorization": f"Bearer {token}"},
    )
    resume_id = upload_res.json()["id"]

    parse_res = client.post(
        "/api/v1/resumes/{}/parse".format(resume_id),
        headers={"Authorization": f"Bearer {token}"},
    )
    assert parse_res.status_code == status.HTTP_200_OK

    get_res = client.get(
        "/api/v1/resumes/{}".format(resume_id),
        headers={"Authorization": f"Bearer {token}"},
    )
    text = get_res.json()["extracted_text"]
    assert "Page 1:" in text
    assert "Page 2:" in text
    assert "Page 3:" in text


def test_parse_docx_with_tables(client):
    """Test extracting text from a DOCX resume containing paragraphs and tables."""
    token = get_auth_token(client, email="docx.parse@example.com")
    docx_bytes = create_mock_docx_bytes()

    upload_res = client.post(
        "/api/v1/resumes/upload",
        files={
            "file": (
                "jane_candidate.docx",
                io.BytesIO(docx_bytes),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers={"Authorization": f"Bearer {token}"},
    )
    resume_id = upload_res.json()["id"]

    parse_res = client.post(
        "/api/v1/resumes/{}/parse".format(resume_id),
        headers={"Authorization": f"Bearer {token}"},
    )
    assert parse_res.status_code == status.HTTP_200_OK

    get_res = client.get(
        "/api/v1/resumes/{}".format(resume_id),
        headers={"Authorization": f"Bearer {token}"},
    )
    text = get_res.json()["extracted_text"]
    assert "Jane Candidate" in text
    assert "Experienced developer proficient" in text
    assert "Skill | Proficiency" in text or "FastAPI | Expert" in text


def test_parse_missing_resume(client):
    """Test parsing a nonexistent resume ID returns 404."""
    token = get_auth_token(client, email="missing@example.com")
    response = client.post(
        "/api/v1/resumes/99999/parse",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == status.HTTP_404_NOT_FOUND
    assert "not found" in response.json()["detail"].lower()


def test_parse_and_get_another_users_resume(client):
    """Test that a user cannot parse or view another user's resume."""
    token1 = get_auth_token(client, email="owner.user@example.com", name="Owner")
    token2 = get_auth_token(client, email="other.user@example.com", name="Other")

    # User 1 uploads
    pdf_bytes = create_mock_pdf_bytes("Confidential Resume for Owner")
    upload_res = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("owner_resume.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
        headers={"Authorization": f"Bearer {token1}"},
    )
    resume_id = upload_res.json()["id"]

    # User 2 attempts to parse User 1's resume
    parse_attempt = client.post(
        "/api/v1/resumes/{}/parse".format(resume_id),
        headers={"Authorization": f"Bearer {token2}"},
    )
    assert parse_attempt.status_code == status.HTTP_404_NOT_FOUND

    # User 2 attempts to fetch User 1's resume
    get_attempt = client.get(
        "/api/v1/resumes/{}".format(resume_id),
        headers={"Authorization": f"Bearer {token2}"},
    )
    assert get_attempt.status_code == status.HTTP_404_NOT_FOUND


def test_parse_corrupted_file(client, db_session):
    """Test parsing when underlying file on disk is corrupted returns 400."""
    token = get_auth_token(client, email="corrupt.parse@example.com")
    pdf_bytes = create_mock_pdf_bytes("Valid initially")

    upload_res = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("corrupt_target.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
        headers={"Authorization": f"Bearer {token}"},
    )
    resume_id = upload_res.json()["id"]

    # Corrupt the file on disk
    resume_db = db_session.query(Resume).filter(Resume.id == resume_id).first()
    with open(resume_db.file_path, "wb") as f:
        f.write(b"CORRUPTED_GARBAGE_DATA_THAT_IS_NOT_A_VALID_PDF")

    parse_res = client.post(
        "/api/v1/resumes/{}/parse".format(resume_id),
        headers={"Authorization": f"Bearer {token}"},
    )
    assert parse_res.status_code == status.HTTP_400_BAD_REQUEST
    assert "corrupted" in parse_res.json()["detail"].lower()
