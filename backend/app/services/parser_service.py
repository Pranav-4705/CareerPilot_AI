import os
import re
from fastapi import HTTPException, status
import fitz  # PyMuPDF
import docx


def clean_extracted_text(text: str) -> str:
    """
    Cleans raw extracted text:
    - Normalizes line breaks
    - Collapses repetitive horizontal whitespace while preserving paragraph structure
    - Limits consecutive newlines to maximum of two
    - Strips leading and trailing whitespace
    """
    if not text:
        return ""

    # Normalize line breaks
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Clean whitespace per line
    lines = []
    for line in text.split("\n"):
        # Collapse multiple spaces/tabs into single space
        cleaned_line = re.sub(r"[ \t]+", " ", line).strip()
        lines.append(cleaned_line)

    cleaned_text = "\n".join(lines)

    # Collapse 3 or more consecutive newlines down to 2
    cleaned_text = re.sub(r"\n{3,}", "\n\n", cleaned_text)

    return cleaned_text.strip()


def extract_text_from_pdf(file_path: str) -> str:
    """
    Extracts text from multi-page PDF documents using PyMuPDF (fitz).
    """
    try:
        doc = fitz.open(file_path)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not open PDF file. File may be corrupted: {str(e)}",
        )

    try:
        pages_text = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text:
                pages_text.append(text)

        full_text = "\n\n".join(pages_text)
        return clean_extracted_text(full_text)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Error extracting text from PDF: {str(e)}",
        )
    finally:
        doc.close()


def extract_text_from_docx(file_path: str) -> str:
    """
    Extracts text from DOCX documents including paragraphs and tables using python-docx.
    """
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not open DOCX file. File may be corrupted: {str(e)}",
        )

    try:
        content_parts = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content_parts.append(text)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                # Deduplicate cells across merged cells in table row
                seen_cells = set()
                row_texts = []
                for cell in row.cells:
                    if cell not in seen_cells and cell.text.strip():
                        seen_cells.add(cell)
                        row_texts.append(cell.text.strip())
                if row_texts:
                    content_parts.append(" | ".join(row_texts))

        full_text = "\n\n".join(content_parts)
        return clean_extracted_text(full_text)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Error extracting text from DOCX: {str(e)}",
        )


def extract_text_from_file(file_path: str) -> str:
    """
    Determines file type and routes to the appropriate parser (PDF or DOCX).
    """
    if not os.path.exists(file_path):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Stored file not found at path: {file_path}",
        )

    _, ext = os.path.splitext(file_path)
    ext_lower = ext.lower()

    if ext_lower == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext_lower == ".docx":
        return extract_text_from_docx(file_path)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Only PDF and DOCX files can be parsed.",
        )
